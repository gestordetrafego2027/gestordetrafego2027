from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.comments import Comment

# shot names from the prompt generator
src=open('/Users/angelomazzutti/Desktop/gestordetrafego2027/docs/shop-confeccao/scripts/gen_prompts.py').read()
ns={}; exec(src.split('def build')[0], ns); PR=ns['P']

F='Arial'
H1=Font(name=F,size=14,bold=True,color='FFFFFF'); FILL_H1=PatternFill('solid',fgColor='1F1F1F')
H2=Font(name=F,size=11,bold=True); FILL_H2=PatternFill('solid',fgColor='D9D9D9')
TH=Font(name=F,size=10,bold=True); FILL_TH=PatternFill('solid',fgColor='F2F2F2')
N=Font(name=F,size=10); INP=Font(name=F,size=10,color='0000FF'); LNK=Font(name=F,size=10,color='008000')
YEL=PatternFill('solid',fgColor='FFFF00'); RED=PatternFill('solid',fgColor='F8D7DA'); GRN=PatternFill('solid',fgColor='D4EDDA')
thin=Side(style='thin',color='BFBFBF'); BR=Border(left=thin,right=thin,top=thin,bottom=thin)
WR=Alignment(wrap_text=True,vertical='top')
BRL='R$ #,##0.00;(R$ #,##0.00);-'; PCT='0.0%'

wb=Workbook()

# ---------- Legenda ----------
lg=wb.active; lg.title='Legenda'
rows=[('Planejamento de Produto e Desenvolvimento — Confecção HMZT',),(),
('Fonte','Prints das páginas da Lunfe Malhas (29–30/09/2026) + textos "Descrição Geral" colados pelo Angelo. Fichas em docs/shop-confeccao/catalogo, dossiês em /dossies, prompts em /prompts.'),
('Abas','Resumo · Regras Fornecedor · uma aba por produto (P01–P29) · Não Catalogados · Pendências Gerais · Prompts Lookbook (um prompt por foto)'),(),
('Convenção de cores',''),
('Texto azul','Valor digitado (dado do fornecedor ou estimativa). Pode ser editado.'),
('Texto preto','Fórmula — não editar.'),
('Texto verde','Fórmula que puxa de outra aba.'),
('Fundo amarelo','Estimativa HMZT ou campo a preencher — revisar quando chegar o orçamento real.'),
('Fundo vermelho claro','Contradição, risco ou dado faltando.'),(),
('Marcas nas tabelas',''),
('✅','Dado técnico declarado pelo fornecedor'),
('📣','Alegação de marketing do fornecedor — NÃO usar como fato na comunicação HMZT'),
('⚠️','Contradição ou erro encontrado'),
('🟢','Recomendação HMZT'),(),
('Como usar','1) Atualize descontos/Pix em "Regras Fornecedor" — todas as abas recalculam. 2) Em cada produto, preencha os custos amarelos (personalização, frete, embalagem, impostos) quando chegarem os orçamentos. 3) Ajuste o preço de venda HMZT e veja a margem.'),
]
for r in rows: lg.append(list(r))
lg['A1'].font=Font(name=F,size=14,bold=True)
for row in lg.iter_rows(min_row=2):
    for c in row: c.font=N; c.alignment=WR
for r in (6,13): lg.cell(r,1).font=H2
for r in range(3,20):
    if lg.cell(r,1).value in ('Fonte','Abas','Como usar'): lg.cell(r,1).font=TH
lg['B7'].font=INP; lg['B8'].font=N; lg['B9'].font=LNK; lg['B10'].fill=YEL; lg['B11'].fill=RED
lg.column_dimensions['A'].width=22; lg.column_dimensions['B'].width=110

# ---------- Regras Fornecedor ----------
rg=wb.create_sheet('Regras Fornecedor')
rg['A1']='Regras gerais — Lunfe Malhas (fabricação própria, Itajaí-SC)'; rg['A1'].font=Font(name=F,size=13,bold=True)
rg.append([]); rg.append(['Faixa','De (pç)','Até (pç)','Desconto'])
tiers=[('Faixa 1',20,40,0.10),('Faixa 2',40,80,0.15),('Faixa 3',80,100,0.20),('Faixa 4',100,300,0.25),('Faixa 5',300,999,0.30)]
for t in tiers: rg.append(list(t))
for c in rg[3]: c.font=TH; c.fill=FILL_TH; c.border=BR
for r in range(4,9):
    for c in range(1,5):
        cell=rg.cell(r,c); cell.border=BR; cell.font=INP if c>1 else N
    rg.cell(r,4).number_format=PCT
rg['A10']='Desconto Pix adicional'; rg['B10']=0.05; rg['B10'].number_format=PCT; rg['B10'].font=INP
rg['A11']='Pedido mínimo (pç, pedido total)'; rg['B11']=20; rg['B11'].font=INP
rg['C10']='Sobre o preço já com desconto progressivo'
other=[('Cartão','Até 6× sem juros no atacado (página do produto diz "com juros" ⚠️ — confirmar)'),
('Boleto','30/60 dias para cadastro aprovado acima de R$ 2.000'),
('Amostra','Até 2 peças, cobradas e abatidas no pedido aprovado; frete por conta do cliente'),
('Mistura','Mínimo de 20 peças no pedido total — mistura livre de produtos, modelos, cores e tamanhos; soma para o desconto'),
('Cor Pantone exclusiva','A partir de 100 peças/cor; 15–25 dias úteis'),
('Embalagem customizada','A partir de 50 peças'),
('Personalização','Silk, DTF, bordado, private label a partir de 20 peças — preço SEMPRE sob orçamento'),
('Prazo','Lisas em estoque: despacho em 24h úteis · personalizadas: sob consulta'),
('Grade customizada','Acima de 200 peças (sob consulta)'),
('Molde regular','Tabela única em todos os regulares: M = 72 × 54 cm, ombro 46, manga 18; P–EG, +2 cm por tamanho'),
('Modelo das fotos','1,84 m · 84 kg · veste M'),
('Nota fiscal','NF-e com CFOP de revenda; Simples, Presumido e Real'),
('Fabricação','Própria: tecelagem, tingimento, corte, confecção e tratamentos em Itajaí-SC')]
r=13; rg.cell(r,1,'Outras condições').font=H2
for k,v in other:
    r+=1; rg.cell(r,1,k).font=TH; c=rg.cell(r,2,v); c.font=N; c.alignment=WR
    rg.merge_cells(start_row=r,start_column=2,end_row=r,end_column=4)
rg.column_dimensions['A'].width=30
for col in 'BCD': rg.column_dimensions[col].width=38

DISC=["'Regras Fornecedor'!$D$%d"%i for i in range(4,9)]
PIX="'Regras Fornecedor'!$B$10"

# ---------- product data ----------
STD=[('P',70,52,44,16),('M',72,54,46,18),('G',74,56,48,20),('GG',76,58,50,22),('EG',78,60,52,24)]
D={}
D['P01']=dict(nome='Camiseta Oversized Suedine Peruano 300g',cod='HMZT-P01-OVS-SUE300',ref='7897385478523 (EAN-13)',
 url='lunfemalhas.com.br/camiseta-oversized-suedine-peruano-300g-atacado-lunfe',cat='Camisetas › Masculino › Oversized',
 base=54.00,linha='Heavy (oversized)',varejo=169,pers=16,frete=6,estoque='Em estoque',
 spec=[('Modelagem','Oversized boxy, ombro caído (drop shoulder)','✅'),('Gramatura','300 g/m² (heavyweight)','✅'),
  ('Tecido','Suedine — acabamento escovado (peach finish), toque de camurça','✅'),('Composição','52% algodão / 48% poliéster','✅'),
  ('Gola','Careca, ribana canelada média','✅'),('Manga','Curta, larga, até acima do cotovelo','✅'),('Barra','Reta, pouco abaixo do quadril','✅'),
  ('Opacidade','Total (mesma malha do P04)','✅'),('Irmão','P04 — mesma malha, modelagem regular (texto do P04 confirma)','✅'),
  ('Cuidados','Lavar do avesso a 30 °C, sem secadora quente, não passar ferro direto na superfície','✅')],
 med=None, med_note='⚠️ Tabela da oversized NÃO informada — pedir ao fornecedor (referência provisória: Chinese Heavyweight oversized, aba P05).',
 cores=[('Off-white','OFF','—','Cor das fotos'),('Areia / bege','ARE','—',''),('Cinza chumbo','CHU','—','')],
 cores_note='⚠️ Cartela da página. O texto do P04 (mesma malha) lista Preto, Off White, Bege, Marrom, Marinho (refs SDP-). Confirmar a cartela da oversized.',
 pers_list=[('Silk base d\'água','Testar aderência sobre suedine','20 pç'),('DTF','Testar aderência','20 pç'),('Bordado','Melhor resultado; costas na oversized','20 pç'),('Etiqueta private label','Interna + tag','20 pç'),('Embalagem customizada','','50 pç'),('Cor Pantone','','100 pç/cor')],
 contra=['"Peruano" no nome, mas é 52/48 com poliéster — não é algodão peruano','Tabela de medidas da oversized não informada','Cartela da página × texto do P04 divergem'],
 pode=['Toque aveludado tipo camurça; acabamento físico que não sai na lavagem','Heavyweight 300 g/m², 100% opaca','Composição por extenso: 52% algodão, 48% poliéster','Fabricada em Santa Catarina'],
 naopode=['"Algodão peruano" (NUNCA)','"Maior conversão por toque", "exclusiva no atacado"','"Percepção de R$ 149–199"'],
 pend=['Tabela de medidas oversized','Cartela real de cores','O que significa "Peruano" no nome','Peça-teste de silk/DTF','Orçamento de personalização'],
 descritor='Camiseta masculina heavyweight oversized boxy, ombro caído, manga larga até acima do cotovelo, gola careca canelada, malha aveludada e fosca tipo camurça.')
D['P02']=dict(nome='Camiseta Oversized Algodão Penteado 185g',cod='HMZT-P02-OVS-ALG185',ref='7899196241006 (EAN-13)',
 url='lunfemalhas.com.br/camiseta-oversized-drop-shoulder-atacado-185g-lunfe',cat='Camisetas › Masculino › Oversized',
 base=39.78,linha='Essencial (alternativa)',varejo=99,pers=9,frete=4,estoque='Disponibilidade imediata',
 spec=[('Modelagem','Oversized "regular drop" — ombro caído, caimento amplo','✅'),('Gramatura','185 g (título/URL) ⚠️ resumo diz 300 g/m²','⚠️'),
  ('Tecido','Malha de algodão penteado','✅'),('Composição','100% algodão penteado','✅'),('Gola','Careca, ribana canelada fina','✅'),
  ('Manga','Curta, larga, até acima do cotovelo','✅'),('Barra','Reta','✅'),('Envio','Até 24h; produção própria Lunfe','✅')],
 med=None, med_note='⚠️ Tabela de medidas não informada — falta a Descrição Geral deste produto.',
 cores=[('Preto','PRT','—','Cor das fotos'),('Branco','BCO','—','')],
 cores_note='A 5ª miniatura (banner "Atacado") sugere cartela maior — confirmar.',
 pers_list=[('Estampa (silk/DTF)','Frente — área mostrada no mockup do fornecedor','20 pç'),('Bordado','','20 pç'),('Etiqueta private label','','20 pç')],
 contra=['Gramatura: 185 g no título × 300 g/m² no resumo','Descrição Geral não enviada — dados técnicos incompletos'],
 pode=['100% algodão penteado','Oversized com ombro caído'],
 naopode=['Gramatura, até confirmar'],
 pend=['Descrição Geral completa','Gramatura real','Tabela de medidas','Cartela completa','Orçamento de estampa/bordado'],
 descritor='Camiseta masculina oversized, ombro caído, manga larga até acima do cotovelo, gola careca fina, algodão penteado fosco.')
D['P03']=dict(nome='Camiseta Pima Cotton 165g',cod='HMZT-P03-REG-PIM165',ref='cms-pimacotton (código interno)',
 url='lunfemalhas.com.br/camiseta-pima-cotton-165g-100-algodao-puro-atacado-lunfe',cat='Camisetas › Masculino › Premium',
 base=63.45,linha='Premium (principal)',varejo=189,pers=12.5,frete=7.5,estoque='8 cores em estoque (página mostra 5 ⚠️)',
 spec=[('Composição','100% algodão Pima (G. barbadense) — sem elastano, sem poliéster','✅'),('Fibra','Extra-longa, staple > 35 mm','✅'),
  ('Construção','"Malha premium" — título do fio não informado','⚠️'),('Gramatura','165 g/m²','✅'),('Modelagem','Regular fit','✅'),
  ('Gola','Careca, ribana reforçada','✅'),('Mangas','Curtas, barra dupla','✅'),('Barra','Dupla, acabamento invisível','✅'),
  ('Costuras','Reforçadas; ombro a ombro com fita interna','✅'),('Tratamentos','Pré-encolhimento industrial; fixação de cor de alta solidez','✅'),
  ('Encolhimento','1–2% nas primeiras lavagens, estabiliza após 2–3 ciclos','✅'),('Cuidados','Lavar a 30 °C, sem secadora quente, pendurar; ferro baixo','✅'),
  ('Produção','Própria Lunfe — Itajaí-SC','✅')],
 med=STD, med_note='Em cm, em repouso, ± 2 cm. Molde regular padrão.',
 cores=[('Preto','PRT','PMA-PRT','Não aparecia no seletor'),('Off White','OFF','PMA-OFW','Não aparecia no seletor'),('Bege','BGE','PMA-BGE','Provável "creme"'),
  ('Capuccino','CAP','PMA-CAP','Exclusiva do Pima; provável "caramelo"'),('Cinza Claro','CZC','PMA-CZC',''),('Marrom','MRR','PMA-MRR','Cor das fotos ("café")'),
  ('Marinho','MAR','PMA-MRN',''),('Verde','VRD','PMA-VRD','Não aparecia no seletor')],
 cores_note='Pantone a partir de 100 pç/cor.',
 pers_list=[('Silk base d\'água','Toque zero','20 pç'),('DTF','Full color','20 pç'),('Bordado','Peito, mangas — o mais indicado','20 pç'),('Etiqueta private label','Interna + tag','20 pç'),('Embalagem customizada','Saco, caixa premium, tag','50 pç'),('Cor Pantone','','100 pç/cor')],
 contra=['"Origem: Pima" — Pima é variedade, não lugar (EUA ou Peru?)','Texto do P08 cita "Pima 200g, 100% peruano, sob consulta" — não bate','Ticket varejo: R$ 149–229 (texto P07) × R$ 179–299 (este texto)','Cores: 5 na página × 8 no texto'],
 pode=['100% algodão Pima — SÓ com comprovação do fio','Fibra extra-longa: toque sedoso','Sem elastano, sem poliéster','Malha leve 165 g/m²','Fabricada em Santa Catarina'],
 naopode=['"A fibra mais nobre do planeta"','"Percepção de R$ 199–299", "maior margem"','"Migração irreversível", LTV 4–6×, recompra'],
 pend=['Origem/variedade do Pima + documento do fio (Supima?)','Título do fio','Estoque de Preto, Off White e Verde','Existe uma "Pima 200g"?','Orçamento bordado + etiqueta + caixa'],
 descritor='Camiseta masculina regular fit, manga ajustada até metade do bíceps, gola careca de ribana fina rente, malha leve de algodão sedosa com brilho natural.')
D['P04']=dict(nome='Camiseta Suedine Peruano 300g Regular Fit',cod='HMZT-P04-REG-SUE300',ref='7891865209942 (EAN-13)',
 url='lunfemalhas.com.br/camiseta-suedine-peruano-300g-toque-aveludado-atacado-lunfe',cat='Algodão Importado (breadcrumb) · também em Masculino › Premium',
 base=52.00,linha='Heavy (principal)',varejo=159,pers=12.5,frete=6,estoque='5 cores em estoque',
 spec=[('Composição','52% algodão / 48% poliéster','✅'),('Acabamento','Suedine (peach/sanded finish) — escovação com micro-abrasivos; físico, permanente','✅'),
  ('Gramatura','300 g/m²','✅'),('Modelagem','Regular fit, caimento estruturado sem elastano','✅'),('Gola','Careca, ribana reforçada','✅'),
  ('Mangas','Curtas, barra dupla','✅'),('Barra','Dupla, acabamento invisível','✅'),('Costuras','Reforçadas; ombro a ombro com fita interna','✅'),
  ('Tratamentos','Pré-encolhimento industrial; fixação de cor','✅'),('Opacidade','Total, inclusive off white','✅'),
  ('Cuidados','Lavar do avesso a 30 °C, sem secadora quente, não passar ferro direto','✅'),('Uso','Meia-estação 15–22 °C; mar–out no Sul/Sudeste','📣')],
 med=STD, med_note='Em cm, em repouso, ± 2 cm.',
 cores=[('Preto','PRT','SDP-PRT',''),('Off White','OFF','SDP-OFW','100% opaco'),('Bege','BGE','SDP-BGE','Provável nude/rosé das fotos'),('Marrom','MRR','SDP-MRR','"A que mais parece camurça"'),('Marinho','MAR','SDP-MRN','')],
 cores_note='⚠️ Página mostrava Areia, Cinza chumbo, Branco, Preto — sem Marrom/Marinho. Confirmar.',
 pers_list=[('Silk base d\'água','Frente, costas, mangas — testar aderência','20 pç'),('DTF','Testar aderência','20 pç'),('Bordado','Peito, mangas, costas — melhor resultado','20 pç'),('Etiqueta private label','Interna + tag','20 pç'),('Embalagem customizada','Saco, caixa premium, tag','50 pç'),('Cor Pantone','','100 pç/cor')],
 contra=['"Peruano" no nome — é 52/48 com poliéster, não é algodão peruano','Cores página × texto','"A partir de" na página × preço fixo no texto','Texturizadas 240g: 52/48 aqui × 51% poliéster/49% algodão no texto do P05','Ticket Chinese HW: R$ 79–129 aqui × drops de R$ 150–250 no texto do P05'],
 pode=['Toque aveludado tipo camurça, permanente','Heavyweight 300 g/m², 100% opaca','Caimento estruturado','Disponível em regular e oversized','Composição por extenso: 52% algodão, 48% poliéster'],
 naopode=['"Algodão peruano" (NUNCA)','"Maior conversão por toque", "ciclo de 5 passos"','"ROI de R$ 6–10 por R$ 1"','"Exclusiva no atacado brasileiro"'],
 pend=['Cartela real (cinza chumbo existe?)','Significado de "Peruano" no nome','Peça-teste silk/DTF','Orçamento bordado + etiqueta + caixa'],
 descritor='Camiseta masculina heavyweight regular fit, manga reta, gola careca canelada, malha encorpada aveludada e fosca tipo camurça, caimento estruturado.')
D['P05']=dict(nome='Camiseta Chinese Heavyweight 300g',cod='HMZT-P05-REG-XJG300 / -OVS-',ref='790836512100 (12 dígitos ⚠️)',
 url='lunfemalhas.com.br/camiseta-chinese-heavyweight-atacado-300g-algodao-importado-lunfe',cat='Camisetas › Masculino › Premium',
 base=52.79,linha='Heavy (alternativa ⚠️ Xinjiang)',varejo=169,pers=15,frete=5,estoque='4 cores em pronta entrega permanente',
 spec=[('Composição','100% algodão importado — origem Xinjiang, China ⚠️','✅'),('Fibra','Longa, staple 28–30 mm','✅'),('Fio','28/1 × 2 cabos (retorcido duplo; equivale a ~14/1)','✅'),
  ('Gramatura','300 g/m² ± 5% ("ultra-heavyweight")','✅'),('Rolo / rendimento','1,80 m tubular aberta · 3,0–3,5 camisetas/m (M)','✅'),
  ('Modelagens','Regular (P–GG) + Oversized (P–G1) — mesma malha','✅'),('Gola','Canelada 2/1 estruturada, fecho em anel','✅'),
  ('Ombros','Reforço ombro a ombro com faixa larga','✅'),('Costuras','Ponto cadeia','✅'),('Barras','Dupla nas mangas e barra','✅'),
  ('Tratamentos','Pré-encolhimento (<3%), tingimento reativo, anti-torção, amaciamento, anti-pilling enzimático','✅'),
  ('Anti-rugas','Não','✅'),('Versão feminina','Não existe (em análise)','✅')],
 med=[('P',70,52,44,16),('M',72,54,46,18),('G',74,56,48,20),('GG',76,58,50,22)],
 med2=[('P',74,57,54,24),('M',76,59,56,25),('G',78,61,58,26),('GG',80,64,60,27),('G1',82,67,62,28)],
 med_note='Regular vai só até GG (sem EG ⚠️). ± 2 cm.',
 cores=[('Preto','PRT','CHW-PRT','Maior giro'),('Cinza','CZA','CHW-CZA',''),('Marinho','MAR','CHW-MRN',''),('Off White','OFF','CHW-OFW','')],
 cores_note='Pantone a partir de 100 pç/cor, 15–25 dias úteis.',
 pers_list=[('Silk','Até 4 cores','20 pç'),('DTF','Full color','20 pç'),('Bordado','Até 15.000 pontos','20 pç'),('Sublimação','⚠️ Não funciona em 100% algodão — não usar','20 pç'),('Etiqueta private label','','20 pç'),('Embalagem customizada','','50 pç'),('Cor Pantone','','100 pç/cor')],
 contra=['Sublimação oferecida em 100% algodão','Chinese Cotton citada como 180g (página P06 diz 160g)','Grade regular só até GG','Oversized sem página/SKU visível','Referência com 12 dígitos'],
 pode=['Heavyweight 300 g/m², fio 2 cabos','100% algodão de fibra longa','Gola estruturada, reforço ombro a ombro','Regular e oversized'],
 naopode=['Origem Xinjiang (NUNCA citar)','"95% das heavyweight nacionais…"','"Não forma pilling", "maior recompra"'],
 pend=['Como pedir a oversized (SKU, fotos)','Regular tem EG?','Orçamento silk/DTF A3 + bordado','Documentação de origem','Resto do FAQ'],
 descritor='Camiseta masculina heavyweight regular, gola careca estruturada, malha de algodão grossa e encorpada com leve brilho acetinado, caimento reto.')
D['P06']=dict(nome='Camiseta Chinese Cotton 160g',cod='HMZT-P06-REG-XJG160',ref='7908365128857 (EAN-13)',
 url='lunfemalhas.com.br/camiseta-chinese-cotton-atacado-lunfe',cat='Camisetas › Masculino › Premium',
 base=44.20,linha='Essencial (alternativa ⚠️ Xinjiang)',varejo=119,pers=9,frete=4,estoque='5 cores em pronta entrega permanente',
 spec=[('Composição','100% algodão importado — Xinjiang, China ⚠️','✅'),('Fibra','Longa, staple 28–30 mm','✅'),('Fio','Penteado, 30/1','✅'),
  ('Gramatura','160 g/m² ± 5% (⚠️ texto do P05 diz 180g)','⚠️'),('Rolo / rendimento','1,80 m tubular aberta · 4,0–4,5 camisetas/m (M)','✅'),
  ('Modelagem','Regular fit','✅'),('Gola','Canelada 2/1, fecho em anel (sem emenda frontal)','✅'),('Ombros','Reforço ombro a ombro','✅'),
  ('Costuras','Ponto cadeia','✅'),('Barras','Dupla nas mangas e barra','✅'),
  ('Tratamentos','Pré-encolhimento (<3%), tingimento reativo, anti-torção, anti-pilling enzimático, amaciamento','✅')],
 med=STD, med_note='± 2 cm. Grade customizada acima de 200 pç.',
 cores=[('Preto','PRT','CHC-PRT','Maior giro'),('Marinho','MAR','CHC-MRN','2º maior giro'),('Verde Militar','VML','CHC-VML',''),('Azul Celeste','AZC','CHC-AZC',''),('Off White','OFF','CHC-OFW','Cor das fotos')],
 cores_note='Pantone a partir de 100 pç/cor.',
 pers_list=[('Silk base d\'água / plastisol / gel / puff 3D / foil / silicone 3D','6 técnicas','20 pç'),('DTF','Full color','20 pç'),('Bordado','📣 "até 15.000 pontos" — preferir ≤ 8.000 em 160g','20 pç'),('Etiqueta private label','','20 pç'),('Embalagem customizada','','50 pç'),('Cor Pantone','','100 pç/cor')],
 contra=['160g × 180g entre textos','Fibra 28–30 mm × 28–32 mm (texto P07)','"Rastreabilidade completa" confirma a origem Xinjiang','Bordado 15.000 pontos em 160g é otimista'],
 pode=['Algodão de fibra longa, fio penteado','Pré-encolhida, anti-torção, tingimento reativo','Fabricada em Santa Catarina'],
 naopode=['Origem Xinjiang (NUNCA citar)','"Campeã de vendas", "maior recompra"','"Amacia com as lavagens"'],
 pend=['Gramatura real (160 × 180)','Orçamento de personalização','Documentação de origem','Comparar com Prime Cotton BCI'],
 descritor='Camiseta masculina regular fit, gola careca canelada fina, malha leve de algodão penteado com leve brilho acetinado.')
D['P07']=dict(nome='Camiseta Cotton Peruano 220g com Elastano',cod='HMZT-P07-REG-PER220',ref='7892474270330 (EAN-13)',
 url='lunfemalhas.com.br/camiseta-cotton-peruano-220g-elastano-atacado-lunfe',cat='Camisetas › Masculino › Premium',
 base=55.00,linha='Premium (alternativa)',varejo=149,pers=11.5,frete=6,estoque='🔴 Página: "Esse acabou" × texto: 6 cores em estoque',
 spec=[('Composição','96% algodão peruano (G. barbadense) + 4% elastano','✅'),('Fibra','Extra-longa, staple > 34 mm (Pima peruano ou Tangüis?)','⚠️'),
  ('Construção','Jersey 30/1','✅'),('Gramatura','220 g/m²','✅'),('Modelagem','Regular fit','✅'),('Gola','Careca, ribana reforçada','✅'),
  ('Mangas','Curtas, barra dupla','✅'),('Barra','Dupla, acabamento invisível','✅'),('Costuras','Reforçadas; ombro a ombro com fita','✅'),
  ('Tratamentos','Pré-encolhimento; fixação de cor de alta solidez','✅'),('Elastano','Memória de forma, caimento, menos rugas','✅'),('Opacidade','Opaca em todas as cores','✅')],
 med=STD, med_note='Em cm, em repouso, ± 2 cm.',
 cores=[('Preto','PRT','CPR-PRT','Cor das fotos'),('Off White','OFF','CPR-OFW',''),('Bege','BGE','CPR-BGE','Provável nude da 5ª foto'),('Marinho','MAR','CPR-MRN',''),('Bordô','BDO','CPR-BDO',''),('Verde Militar','VML','CPR-VML','')],
 cores_note='Pantone a partir de 100 pç/cor, 15–25 dias úteis.',
 pers_list=[('Silk base d\'água','Toque zero','20 pç'),('DTF','Full color','20 pç'),('Bordado','Peito, mangas','20 pç'),('Sublimação','⚠️ Não funciona com 96% algodão — não usar','20 pç'),('Etiqueta private label','Interna + tag','20 pç'),('Embalagem customizada','','50 pç'),('Cor Pantone','','100 pç/cor')],
 contra=['Esgotado na página × em estoque no texto','6× com juros (página) × sem juros (texto)','"30/1s mais grosso que 30/1" — mesmo fio','Sublimação oferecida em algodão','Cores das fotos × lista'],
 pode=['Algodão peruano de fibra extra-longa — só com a variedade confirmada','96% algodão, 4% elastano: volta à forma','220 g/m², opaca, pré-encolhida'],
 naopode=['"Migração irreversível", recompra, LTV 3–5×','"Camiseta de R$ 150"','"Aristocracia do algodão"'],
 pend=['Variedade: Pima peruano ou Tangüis?','Reposição','Juros no cartão','Orçamento bordado + etiqueta','Frete Itajaí → SP'],
 descritor='Camiseta masculina regular fit, gola careca de ribana fina, jersey de algodão com leve brilho acetinado que acompanha o corpo.')
D['P08']=dict(nome='Camiseta Egyptian Cotton Fio 50/1 com Elastano',cod='HMZT-P08-REG-EGY50',ref='7891557272995 (EAN-13)',
 url='lunfemalhas.com.br/camiseta-egyptian-cotton-atacado-algodao-egipcio-50-1-lunfe',cat='Camisetas › Masculino › Premium',
 base=52.65,linha='Topo / Assinatura (inclui unissex)',varejo=189,pers=15,frete=6,estoque='7 cores em estoque (off white com traço na página ⚠️)',
 spec=[('Composição','91,5% algodão egípcio (G. barbadense) + 8,5% elastano','✅'),('Fibra','Extra-longa, staple 32–36 mm','✅'),
  ('Fio','Penteado e refinado 50/1 (o mais fino do catálogo)','✅'),('Gramatura','170 g/m²','✅'),('Modelagem','Regular fit, caimento fluido','✅'),
  ('Gola','Ribana 2/1 lisa (sem canelado)','✅'),('Costura interna','Ponto conjugado (costura plana)','✅'),('Ombros','Reforço ombro a ombro','✅'),
  ('Barra','Dupla','✅'),('Tratamentos','Pré-encolhimento, fixação reativa de cor, amaciamento','✅'),
  ('Certificação','OEKO-TEX Standard 100 (algodão e elastano)','✅')],
 med=STD, med_note='Em cm, em repouso, ± 2 cm.',
 cores=[('Off White','OFF','EGY-OFW','Traço na bolinha — confirmar estoque'),('Bordô','BDO','EGY-BRD',''),('Bege','BGE','EGY-BGE','7ª cor que faltava no seletor'),('Preto','PRT','EGY-PRT',''),('Verde Militar','VML','EGY-VML',''),('Laranja','LAR','EGY-LRJ',''),('Verde','VRD','EGY-VRD','')],
 cores_note='Pantone a partir de 100 pç/cor, 15–25 dias úteis.',
 pers_list=[('Silk base d\'água','Toque zero (evitar plastisol)','20 pç'),('DTF','Full color','20 pç'),('Bordado','Exige entretela (elastano)','20 pç'),('Etiqueta private label','Interna + tag','20 pç'),('Embalagem customizada','','50 pç'),('Cor Pantone','','100 pç/cor')],
 contra=['"50/1 é quase o dobro da finura do 30/1" — exagero','"Como percal 400 fios" — sem base','Fibra 32–36 × 36+ mm entre textos','Tabela cita "Pima 200g peruano sob consulta"','Off white com traço × em estoque'],
 pode=['OEKO-TEX Standard 100 — com nº do certificado','Algodão egípcio fibra extra-longa, fio 50/1 — só com comprovação','Caimento fluido, não amassa','Costura plana, gola lisa'],
 naopode=['"O algodão mais nobre do mundo", "como seda"','"Maior conversão por toque"'],
 pend=['Nº do certificado OEKO-TEX','Comprovação egípcio (Giza? Cotton Egypt?)','Estoque off white','Modelagem feminina?','Orçamento bordado com entretela'],
 descritor='Camiseta regular fit de caimento fluido, gola careca de ribana lisa fina, malha leve de algodão muito lisa com brilho natural sutil.')
D['P09']=dict(nome='Camiseta Suedine Light',cod='HMZT-P09-REG-SUELT',ref='7908365128002 (EAN-13)',
 url='lunfemalhas.com.br/camiseta-suedine-light-atacado-toque-aveludado-lunfe',cat='Camisetas › Masculino › Premium',
 base=44.80,linha='Suedine light',varejo=129,pers=10,frete=4,estoque='Lisas em estoque, 24h',
 spec=[('Composição','Blend algodão + fibras sintéticas — % NÃO informado','⚠️'),('Gramatura','"Leve" — número NÃO informado','⚠️'),
  ('Acabamento','Suedine aveludado','✅'),('Modelagem','Regular fit, caimento fluido','✅'),('Gola','Careca','✅'),
  ('Acabamentos','Costuras reforçadas, pré-encolhida','✅'),('Uso','Ano inteiro; segunda pele sob jaqueta','📣')],
 med=STD, med_note='Em cm, variação 1–2 cm.',
 cores=[('Bege','BGE','—','⚠️ Fotos em caramelo — é este?'),('Preto','PRT','—',''),('Azul Bebê','AZB','—',''),('Laranja','LAR','—',''),('Verde','VRD','—','')],
 cores_note='Mockup do fornecedor mostra off white (fora da lista). Sem códigos de referência.',
 pers_list=[('Silk','Testar aderência','20 pç?'),('DTF','Testar aderência','20 pç?'),('Bordado','','20 pç?'),('Etiqueta private label','','20 pç?')],
 contra=['"Textura que não existe em outra composição" — P01/P04 também são suedine','FAQ diz que a Suedine 300g é oversized (existe a regular)','Caramelo nas fotos × Bege na lista','Sem gramatura, % e condições de pagamento'],
 pode=['Toque aveludado de suedine em malha leve','Pré-encolhida','Composição exata só depois de informada'],
 naopode=['"Premium antes de olhar a etiqueta"','"Não bole"'],
 pend=['Gramatura em g/m²','Composição em %','Referências de cor / caramelo','Condições de pagamento','Mínimos e preço de personalização'],
 descritor='Camiseta masculina regular fit, gola careca canelada, malha leve aveludada e fosca, caimento reto e fluido.')

# ---------- produtos vindos de JSON (P02 revisado, P10+) ----------
import json,glob,os
DATA=os.path.join(os.path.dirname(os.path.abspath(__file__)),'data')
JMETA={}
for f in sorted(glob.glob(os.path.join(DATA,'P*.json'))):
    j=json.load(open(f,encoding='utf-8')); pk=j['pk']
    D[pk]={k:j.get(k) for k in ['nome','cod','ref','url','cat','base','linha','varejo','pers','frete','estoque','spec','med','med2','med_note','cores','cores_note','pers_list','contra','pode','naopode','pend','descritor']}
    D[pk]['spec']=[tuple(x) for x in j['spec']]
    D[pk]['med']=[tuple(x) for x in j['med']] if j.get('med') else None
    D[pk]['med2']=[tuple(x) for x in j['med2']] if j.get('med2') else None
    D[pk]['med_head']=j.get('med_head'); D[pk]['med2_title']=j.get('med2_title')
    D[pk]['cores']=[tuple(x) for x in j['cores']]; D[pk]['pers_list']=[tuple(x) for x in j['pers_list']]
    D[pk]['categoria']=j.get('categoria_produto','Camiseta')
    def find(*keys):
        for it,val,_ in j['spec']:
            if any(k.lower() in it.lower() for k in keys): return val
        return '—'
    JMETA[pk]=(find('Modelagem','Corte'),find('Tecido','Malha','Construção'),find('Gramatura'),find('Composição'),(j['contra'][0] if j.get('contra') else '—'))
D=dict(sorted(D.items()))
for pk in D: D[pk].setdefault('categoria','Camiseta')

# ---------- product sheets ----------
def hdr(ws,r,text,cols=6):
    ws.cell(r,1,text).font=H2
    for c in range(1,cols+1): ws.cell(r,c).fill=FILL_H2
    return r+1
def table(ws,r,head,rows,fonts=None,fills=None):
    for i,h in enumerate(head,1):
        c=ws.cell(r,i,h); c.font=TH; c.fill=FILL_TH; c.border=BR; c.alignment=WR
    r+=1
    for row in rows:
        for i,v in enumerate(row,1):
            c=ws.cell(r,i,v); c.font=INP if (fonts and fonts[i-1]=='i') else N; c.border=BR; c.alignment=WR
            if isinstance(v,str) and v.startswith('⚠️'): c.fill=RED
        r+=1
    return r
REFS={}
for pk,d in D.items():
    ws=wb.create_sheet(pk)
    ws.cell(1,1,f"{pk} — {d['nome']}").font=H1
    for c in range(1,7): ws.cell(1,c).fill=FILL_H1
    r=3; r=hdr(ws,r,'1. Identificação')
    ident=[('Código HMZT',d['cod']),('Fornecedor','Lunfe Malhas — fabricação própria, Itajaí-SC'),('Referência',d['ref']),('URL',d['url']),
           ('Categoria no site',d['cat']),('Linha HMZT recomendada',d['linha']),('Estoque',d['estoque']),('Descritor canônico (IA)',d['descritor'])]
    for k,v in ident:
        ws.cell(r,1,k).font=TH; c=ws.cell(r,2,v); c.font=N; c.alignment=WR; ws.merge_cells(start_row=r,start_column=2,end_row=r,end_column=6)
        if '⚠️' in v or '🔴' in v: c.fill=RED
        r+=1
    r+=1; r=hdr(ws,r,'2. Especificação técnica')
    r=table(ws,r,['Item','Valor','Fonte'],[(a,b,c) for a,b,c in d['spec']])
    for rr in range(r-len(d['spec']),r): ws.merge_cells(start_row=rr,start_column=2,end_row=rr,end_column=5); ws.cell(rr,6,ws.cell(rr,3).value) if False else None
    r+=1; r=hdr(ws,r,'3. Tabela de medidas (cm)')
    if d.get('med'):
        mh=d.get('med_head') or ['Tamanho','Comprimento','Largura','Ombros','Mangas']
        if d.get('med2'): ws.cell(r,1,'Regular' if not d.get('med2_title') else 'Tabela 1').font=TH; r+=1
        r=table(ws,r,mh,d['med'],fonts=['n']+['i']*(len(mh)-1))
        if d.get('med2'):
            ws.cell(r,1,d.get('med2_title') or 'Oversized').font=TH; r+=1
            r=table(ws,r,mh,d['med2'],fonts=['n']+['i']*(len(mh)-1))
    c=ws.cell(r,1,d['med_note']); c.font=N; c.alignment=WR; ws.merge_cells(start_row=r,start_column=1,end_row=r,end_column=6)
    if d['med_note'].startswith('⚠️'): c.fill=RED
    ws.row_dimensions[r].height=30; r+=1
    ws.cell(r,1,'Modelo das fotos: 1,84 m · 84 kg · veste M').font=N; r+=2
    r=hdr(ws,r,'4. Cores disponíveis')
    r=table(ws,r,['Cor','Código HMZT','Ref. fornecedor','Observação'],d['cores'])
    c=ws.cell(r,1,d['cores_note']); c.font=N; c.alignment=WR; ws.merge_cells(start_row=r,start_column=1,end_row=r,end_column=6)
    if d['cores_note'].startswith('⚠️'): c.fill=RED
    ws.row_dimensions[r].height=30; r+=2
    r=hdr(ws,r,'5. Personalização (preço sempre sob orçamento)')
    r=table(ws,r,['Técnica','Detalhe','Mínimo'],d['pers_list']); r+=1
    # prices
    r=hdr(ws,r,'6. Preços do fornecedor (desconto progressivo)')
    ws.cell(r,1,'Preço base (R$)').font=TH; b=ws.cell(r,2,d['base']); b.font=INP; b.number_format=BRL
    base_ref=f"$B${r}"
    if pk=='P07': b.comment=Comment('Confirmado no texto da Descrição Geral (a página estava esgotada).','Claude')
    r+=1
    for i,h in enumerate(['Faixa','Peças','Desconto','Preço unitário','Preço + Pix'],1):
        c=ws.cell(r,i,h); c.font=TH; c.fill=FILL_TH; c.border=BR
    r+=1; first=r
    for i,(nm,de,ate,_) in enumerate(tiers):
        ws.cell(r,1,nm).font=N
        ws.cell(r,2,f"{de}–{ate}").font=N
        c=ws.cell(r,3,f"={DISC[i]}"); c.font=LNK; c.number_format=PCT
        c=ws.cell(r,4,f"=ROUND({base_ref}*(1-C{r}),2)"); c.number_format=BRL; c.font=N
        c=ws.cell(r,5,f"=ROUND(D{r}*(1-{PIX}),2)"); c.number_format=BRL; c.font=LNK
        for cc in range(1,6): ws.cell(r,cc).border=BR
        r+=1
    lote20=f"E{first}"
    r+=1
    # cost & margin
    r=hdr(ws,r,'7. Custo HMZT e margem (lote de 20 peças, Pix)')
    rows=[('Peça (20 pç, Pix)',f"={lote20}",'f',''),
          ('Personalização + etiqueta + tag (R$/pç)',d['pers'],'y','Estimativa HMZT — trocar pelo orçamento real'),
          ('Frete rateado + embalagem (R$/pç)',d['frete'],'y','Estimativa HMZT — trocar pelo orçamento real')]
    start=r
    for k,v,t,note in rows:
        ws.cell(r,1,k).font=TH; c=ws.cell(r,2,v); c.number_format=BRL
        c.font=N if t=='f' else INP
        if t=='y': c.fill=YEL; ws.cell(r,3,note).font=N
        r+=1
    ws.cell(r,1,'Custo unitário estimado').font=TH; c=ws.cell(r,2,f"=SUM(B{start}:B{r-1})"); c.number_format=BRL; c.font=Font(name=F,size=10,bold=True); cost=f"B{r}"; r+=1
    ws.cell(r,1,'Preço de venda HMZT (R$)').font=TH; c=ws.cell(r,2,d['varejo']); c.font=INP; c.fill=YEL; c.number_format=BRL; ws.cell(r,3,'Sugestão do dossiê — ajustar').font=N; price=f"B{r}"; r+=1
    ws.cell(r,1,'Impostos + gateway (% do preço)').font=TH; c=ws.cell(r,2,0); c.font=INP; c.fill=YEL; c.number_format=PCT; ws.cell(r,3,'A preencher (Simples + taxa Asaas)').font=N; tax=f"B{r}"; r+=1
    ws.cell(r,1,'Margem bruta (R$)').font=TH; c=ws.cell(r,2,f"={price}-{cost}"); c.number_format=BRL; c.font=N; r+=1
    ws.cell(r,1,'Margem bruta (%)').font=TH; c=ws.cell(r,2,f"=IF({price}=0,0,({price}-{cost})/{price})"); c.number_format=PCT; c.font=N; r+=1
    ws.cell(r,1,'Margem após impostos/gateway (%)').font=TH; c=ws.cell(r,2,f"=IF({price}=0,0,({price}*(1-{tax})-{cost})/{price})"); c.number_format=PCT; c.font=N
    ws.cell(r,3,'Ainda sem frete ao cliente e mídia').font=N; r+=1
    REFS[pk]=dict(base=base_ref,lote=lote20,top=f"D{first+4}",cost=cost,price=price,mpct=f"B{r-2}")
    r+=1
    r=hdr(ws,r,'8. Contradições e erros ⚠️')
    for x in d['contra']:
        c=ws.cell(r,1,'⚠️ '+x); c.font=N; c.fill=RED; c.alignment=WR; ws.merge_cells(start_row=r,start_column=1,end_row=r,end_column=6); r+=1
    r+=1; r=hdr(ws,r,'9. Comunicação HMZT')
    ws.cell(r,1,'✅ Pode usar').font=TH; ws.cell(r,4,'📣 / ⛔ Não usar').font=TH; r+=1
    n=max(len(d['pode']),len(d['naopode']))
    for i in range(n):
        if i<len(d['pode']):
            c=ws.cell(r,1,d['pode'][i]); c.font=N; c.fill=GRN; c.alignment=WR; ws.merge_cells(start_row=r,start_column=1,end_row=r,end_column=3)
        if i<len(d['naopode']):
            c=ws.cell(r,4,d['naopode'][i]); c.font=N; c.fill=RED; c.alignment=WR; ws.merge_cells(start_row=r,start_column=4,end_row=r,end_column=6)
        r+=1
    r+=1; r=hdr(ws,r,'10. Pendências com o fornecedor')
    r=table(ws,r,['Pendência','Status','Responsável','Resposta'],[(p,'Aberta','','') for p in d['pend']],fonts=['n','i','i','i'])
    r+=1; r=hdr(ws,r,f'11. Fluxo de prompts de lookbook (docs/shop-confeccao/prompts/{pk}-prompts.md)')
    jf=os.path.join(DATA,pk+'.json')
    shots=[tuple(x) for x in json.load(open(jf,encoding='utf-8'))['prompt']['shots']] if os.path.exists(jf) else PR[pk]['shots']+PR[pk].get('extra_variants',[])
    r=table(ws,r,['Foto','Origem no print','Formato','Status'],[(s[0],s[1],s[3],'A gerar') for s in shots],fonts=['n','n','n','i'])
    for col,w in zip('ABCDEF',[34,26,22,22,18,14]): ws.column_dimensions[col].width=w
    ws.freeze_panes='A2'

# ---------- Resumo ----------
rs=wb.create_sheet('Resumo',1)
rs['A1']=f'Resumo do catálogo — {len(D)} produtos Lunfe'; rs['A1'].font=Font(name=F,size=14,bold=True)
head=['Produto','Categoria','Código HMZT','Nome','Linha HMZT','Modelagem','Tecido','Gramatura','Composição','Nº cores','Preço base','Lote 20 + Pix','Topo (300+)','Custo HMZT est.','Preço venda HMZT','Margem bruta %','Risco / atenção']
for i,h in enumerate(head,1):
    c=rs.cell(3,i,h); c.font=TH; c.fill=FILL_TH; c.border=BR; c.alignment=WR
meta={'P01':('Oversized','Suedine','300 g/m²','52% CO / 48% PES','"Peruano" no nome não é algodão peruano'),
 'P02':('Oversized','Algodão penteado','185 g ⚠️','100% CO','Gramatura em conflito; sem Descrição Geral'),
 'P03':('Regular','Pima','165 g/m²','100% Pima','Comprovar Pima'),
 'P04':('Regular','Suedine','300 g/m²','52% CO / 48% PES','"Peruano" no nome; cartela divergente'),
 'P05':('Regular + Oversized','Algodão Xinjiang 2 cabos','300 g/m²','100% CO','Origem Xinjiang'),
 'P06':('Regular','Algodão Xinjiang','160 g/m² ⚠️','100% CO','Origem Xinjiang; 160 × 180 g'),
 'P07':('Regular','Algodão peruano + elastano','220 g/m²','96% CO / 4% EL','Esgotado na página; variedade do algodão'),
 'P08':('Regular','Algodão egípcio 50/1 + elastano','170 g/m²','91,5% CO / 8,5% EL','Comprovar egípcio; nº OEKO-TEX'),
 'P09':('Regular','Suedine light','não informada ⚠️','algodão + sintético (% ?)','Faltam gramatura e composição')}
r=4
for pk,d in D.items():
    m=meta.get(pk) if pk in meta and pk not in JMETA else JMETA[pk]; R=REFS[pk]; q=f"{pk}!"
    vals=[pk,d['categoria'],d['cod'],d['nome'],d['linha'],m[0],m[1],m[2],m[3],len(d['cores']),
          f"={q}{R['base']}",f"={q}{R['lote']}",f"={q}{R['top']}",f"={q}{R['cost']}",f"={q}{R['price']}",f"={q}{R['mpct']}",m[4]]
    for i,v in enumerate(vals,1):
        c=rs.cell(r,i,v); c.border=BR; c.alignment=WR
        c.font=LNK if isinstance(v,str) and v.startswith('=') else (INP if i==10 else N)
        if 11<=i<=15: c.number_format=BRL
        if i==16: c.number_format=PCT
        if i==17: c.fill=RED
    rs.cell(r,1).hyperlink=f"#'{pk}'!A1"
    r+=1
for col,w in zip('ABCDEFGHIJKLMNOPQ',[8,12,24,34,24,16,24,14,18,8,12,12,12,13,13,11,34]): rs.column_dimensions[col].width=w
rs.freeze_panes='C4'
rs.cell(r+1,1,'Valores em verde vêm das abas de cada produto. Custo e preço de venda são estimativas até chegarem os orçamentos.').font=N

# ---------- Não catalogados ----------
nc=wb.create_sheet('Não Catalogados')
nc['A1']='Produtos citados nos textos do fornecedor e sem página no site'; nc['A1'].font=Font(name=F,size=13,bold=True)
rows=[('Prime Cotton BCI 160g','100% algodão nacional BCI (G. hirsutum)','160 g','Fio 30/1 · fibra 24–26 mm (⚠️ texto P07: 24–28)','BCI + OEKO-TEX',36.89,'4 cores','⚠️ Citada nos comparativos, mas NÃO aparece no site (busca e categorias em 30/09/2026) — perguntar se saiu de linha'),
 ('Chinese Heavyweight Oversized','Mesma malha do P05','300 g','Oversized P–G1 (tabela no dossiê P05)','—',52.79,'4 (P05)','Variante do P05 sem página própria'),
 ('Pima Cotton 200g (?)','"100% algodão peruano" (tabela do texto P08)','200 g','—','—',None,'?','Pode ser tabela desatualizada — confirmar'),
 ('"Oversized Suedine 300g" 100% algodão','⚠️ Texto do P02 diz 100% algodão penteado, 6 cores, R$ 42,00','300 g','—','—',42.00,'6','⚠️ Contradiz a ficha P01 (52/48, R$ 54,00) — provável texto desatualizado')]
head=['Produto','Composição','Gramatura','Construção','Certificação','Preço base (R$)','Cores','Observação']
for i,h in enumerate(head,1):
    c=nc.cell(3,i,h); c.font=TH; c.fill=FILL_TH; c.border=BR
for j,row in enumerate(rows,4):
    for i,v in enumerate(row,1):
        c=nc.cell(j,i,v if v is not None else 'sob consulta'); c.border=BR; c.alignment=WR; c.font=INP if i==6 else N
        if i==6 and v is not None: c.number_format=BRL
        if isinstance(v,str) and '⚠️' in v: c.fill=RED
for col,w in zip('ABCDEFGH',[30,40,11,34,16,14,10,40]): nc.column_dimensions[col].width=w

# ---------- Pendências gerais ----------
pg=wb.create_sheet('Pendências Gerais')
pg['A1']='Pendências com o fornecedor — todos os produtos'; pg['A1'].font=Font(name=F,size=13,bold=True)
for i,h in enumerate(['Produto','Pendência','Prioridade','Status','Resposta'],1):
    c=pg.cell(3,i,h); c.font=TH; c.fill=FILL_TH; c.border=BR
r=4
gen=['Orçamento de personalização (silk, DTF, bordado, etiqueta, tag, caixa) por posição','Frete Itajaí → São Paulo para 20 e 50 peças','Tabela de medidas de todas as oversized']
for p in gen:
    for i,v in enumerate(['Geral',p,'Alta','Aberta',''],1):
        c=pg.cell(r,i,v); c.border=BR; c.alignment=WR; c.font=INP if i in (3,4,5) else N
    r+=1
for pk,d in D.items():
    for p in d['pend']:
        pri='Alta' if any(k in p for k in ('Pima','egípcio','OEKO','Gramatura','Composição','Variedade','Origem','Peruano')) else 'Média'
        for i,v in enumerate([pk,p,pri,'Aberta',''],1):
            c=pg.cell(r,i,v); c.border=BR; c.alignment=WR; c.font=INP if i in (3,4,5) else N
        r+=1
for col,w in zip('ABCDE',[10,60,11,11,40]): pg.column_dimensions[col].width=w
pg.freeze_panes='A4'

# ---------- Prompts ----------
import re
pp=wb.create_sheet('Prompts Lookbook')
pp['A1']='Prompts de lookbook — um por foto (colar em inglês no Magnific, modelo ∞, AI prompt desligado)'; pp['A1'].font=Font(name=F,size=13,bold=True)
for i,h in enumerate(['Produto','Foto','Origem no print','Modelo · cor · formato','Prompt (EN)','Status'],1):
    c=pp.cell(3,i,h); c.font=TH; c.fill=FILL_TH; c.border=BR
r=4
PD='/Users/angelomazzutti/Desktop/gestordetrafego2027/docs/shop-confeccao/prompts'
for pk in D:
    txt=open(f'{PD}/{pk}-prompts.md').read()
    for blk in txt.split('\n## ')[1:]:
        name=blk.split('\n')[0]
        orig=re.search(r'Origem no print:\*\* (.*)',blk).group(1)
        meta=re.search(r'- \*\*Modelo:\*\* (.*)',blk).group(1).replace('**','')
        prompt=re.search(r'```text\n(.*?)\n```',blk,re.S).group(1)
        for i,v in enumerate([pk,name,orig,meta,prompt,'A gerar'],1):
            c=pp.cell(r,i,v); c.border=BR; c.alignment=WR; c.font=INP if i==6 else N
        pp.row_dimensions[r].height=300
        r+=1
for col,w in zip('ABCDEF',[8,26,26,22,120,10]): pp.column_dimensions[col].width=w
pp.freeze_panes='A4'
from openpyxl.workbook.properties import CalcProperties
wb.calculation=CalcProperties(fullCalcOnLoad=True)

out='/Users/angelomazzutti/Desktop/gestordetrafego2027/docs/shop-confeccao/Planejamento_Produto_Confeccao_HMZT.xlsx'
wb.save(out); print(out)
